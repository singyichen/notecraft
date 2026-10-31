#!/usr/bin/env python3
"""產生實驗課當天用的資料記錄簿 measured/Lab3-數據記錄.xlsx。

助教公告（11/7 六）只要求實驗一與實驗二：各至少量一條導通的曲線、填表並畫出
I_D-V_DS 與 I_D-V_GS 關係圖；實驗三、四與第二部分行有餘力再做。所以這個活頁簿
只做這兩張表，另附一張照講義版面排好、可直接貼進報告的「交報告用」分頁。

預報欄直接讀 LTspice 的 .raw，所以和筆記裡的數字同源；改了網表重跑本腳本即可。
黃底 = 當天要填的格子，其他欄位都是公式或預報值，填完數據圖就會自己更新。

用法：cd simulations/lab3 && ../lab1/.venv/bin/python build_record_xlsx.py
      （檔案已存在時不覆蓋，加 --force 才重產）
"""
import os
import sys

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.marker import Marker
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', 'tools')))
from ltraw import read_raw  # noqa: E402

OUT = os.path.join(HERE, 'measured', 'Lab3-數據記錄.xlsx')

FONT = 'Arial'
FILL_IN = PatternFill('solid', fgColor='FFF2CC')     # 當天要填
FILL_HD = PatternFill('solid', fgColor='D9E2F3')     # 表頭
FILL_GP = PatternFill('solid', fgColor='EDF1F9')     # 欄組標題
BLUE = Font(name=FONT, color='0000FF')               # 寫死的輸入（模擬值）
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
NOTE = Font(name=FONT, size=9, color='808080')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MARK = '1F4E79'   # 實測點的顏色；點要明確給填色與邊線，否則 LibreOffice 重算存檔會把它寫成 noFill 而隱形

# 講義第 7 頁／第 9 頁的表格級距（前兩列是講義印好的，後面幾列講義留白）
VGS_GROUPS = [0.0, 2.0, 4.0, 6.0]                    # 實驗一：四組 V_GS
VDS_ROWS = [1.0, 2.0, 3.0, 4.0, 5.0]                 # 實驗一：每組五列 V_DS
VDS_GROUPS = [1.0, 3.0, 5.0]                         # 實驗二：三組 V_DS
VGS_ROWS = [0.5, 1.0, 2.0, 3.0, 4.0, 5.0]            # 實驗二：每組六列 V_GS（粗掃，先找抬頭位置）
FINE_N = 10                                          # 加密掃描的列數（找 V_th 用）
R1_EXP1 = 1000.0                                     # 源極取樣電阻
R1_EXP2 = 100.0                                      # 汲極取樣電阻


def dot_series(ys, xs, title):
    """只畫點、不畫線的實測序列。"""
    s = Series(ys, xs, title=title)
    s.marker = Marker(symbol='circle', size=7)
    s.marker.graphicalProperties.solidFill = MARK
    s.marker.graphicalProperties.line.solidFill = MARK
    s.graphicalProperties.line.noFill = True
    s.smooth = False
    return s


def line_series(ys, xs, title):
    # smooth=False：LibreOffice 預設會把折線畫成平滑曲線，V_GS = 6 V 那條會在 4.5 V
    # 附近出現不存在的凸起與回落，看起來像量錯。
    s = Series(ys, xs, title=title)
    s.marker = Marker(symbol='none')
    s.smooth = False
    return s


def put(ws, row, col, value, *, font=BLACK, fill=None, fmt=None, align=None):
    # 黃底輸入格若尚未量測，明確存成空字串型儲存格。
    # 某些 Excel 相容預覽器會把「數值型但沒有值」的儲存格誤顯示成 0。
    if value is None and fill is FILL_IN:
        value = ''
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    c.border = BOX
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = Alignment(horizontal=align, vertical='center')
    return c


def title_block(ws, title, subtitle, span):
    ws['A1'] = title
    ws['A1'].font = TITLE
    ws['A2'] = subtitle
    ws['A2'].font = NOTE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)


def group_head(ws, row, col, span, text):
    """欄組標題（例如「V_GS = 4 V」），橫跨該組所有欄。"""
    c = ws.cell(row=row, column=col, value=text)
    c.font = BOLD
    c.fill = FILL_GP
    c.border = BOX
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
    for k in range(1, span):
        ws.cell(row=row, column=col + k).border = BOX


def col_head(ws, row, col, text, width=None):
    c = ws.cell(row=row, column=col, value=text)
    c.font = BOLD
    c.fill = FILL_HD
    c.border = BOX
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    if width:
        ws.column_dimensions[get_column_letter(col)].width = width
    return c


# ── 讀 LTspice 預報值 ────────────────────────────────────────────────────────
_, E1 = read_raw(os.path.join(HERE, 'lab3-exp1-curves.raw'))
PRED1 = []                                            # [組][列] → I_D (mA)
for s in E1:
    v, idd = s['VDS'], -s['I(VDS)'] * 1e3
    PRED1.append([float(np.interp(x, v, idd)) for x in VDS_ROWS])

_, E2 = read_raw(os.path.join(HERE, 'lab3-exp2-transfer.raw'))
PRED2 = []
for s in E2:
    v, idd = s['VGS'], -s['I(VDS)'] * 1e3
    PRED2.append([float(np.interp(x, v, idd)) for x in VGS_ROWS])

wb = Workbook()

# ── 說明 ────────────────────────────────────────────────────────────────────
ws = wb.active
ws.title = '說明'
title_block(ws, 'Lab 3 MOSFET 實驗　數據記錄簿',
            '電子學實作系列第 9 週（11/7 六）．App_Lab3.pdf．當天用筆電直接填', 6)
rows = [
    ('助教公告的當天範圍', ''),
    ('', '實驗一 I_D-V_DS：至少量一條導通的（V_GS = 2／4／6 V 擇一以上），填表並畫 I_D-V_DS 關係圖'),
    ('', '實驗二 I_D-V_GS：至少量一條導通的（V_DS = 3／5 V 擇一以上），填表並畫 I_D-V_GS 關係圖'),
    ('', '實驗三（導通電阻）、實驗四（LED 開關）、第二部分（共源級放大器）行有餘力再完成'),
    ('', ''),
    ('怎麼用這個檔', ''),
    ('', '黃底 = 當天要填的格子，其他欄位是公式或 LTspice 預報值，不用動'),
    ('', '實驗一填「R1 壓降 V_R1」，I_D 欄會自動算成 V_R1 / 1 kΩ'),
    ('', '實驗二填「汲極對地電壓 V_D」，I_D 欄會自動算成 (V_DS − V_D) / 100 Ω'),
    ('', '兩張散佈圖會跟著填的格子自動長出來，收工前看一眼形狀對不對'),
    ('', '「交報告用」分頁已照講義版面排好，數字由前兩張表帶過來，直接複製貼進報告'),
    ('', ''),
    ('量測注意', ''),
    ('', '實驗一的 R1 串在「源極」：電源讀數 V_GS、V_DS 都含 R1 壓降，不等於元件端點電壓'),
    ('', '實驗二的 R1 串在「汲極」、源極接地：電源讀數 V_GS 就是元件真正的 V_GS'),
    ('', '2N7000 規格書的 V_GS(th) 範圍是 0.8～3.0 V，量到 1.5 V 或 2.8 V 都算正常'),
    ('', 'LTspice 預報用的是 BS170（同族替身），導通門檻與手上那顆不會完全一樣，只拿來對形狀'),
]
for i, (a, b) in enumerate(rows, start=4):
    ca = ws.cell(row=i, column=1, value=a)
    ca.font = BOLD if a else BLACK
    ws.cell(row=i, column=2, value=b).font = BLACK
ws.column_dimensions['A'].width = 22
ws.column_dimensions['B'].width = 96
ws.sheet_view.showGridLines = False


# ── 實驗一：輸出特性曲線 ─────────────────────────────────────────────────────
ws1 = wb.create_sheet('實驗一 特性曲線')
title_block(ws1, '實驗一　輸出特性曲線 I_D-V_DS（講義第 6–7 頁）',
            'R1 = 1 kΩ 串在源極．填 R1 壓降，I_D 自動換算．V_GS = 0 V 那組全是 0 是正常的（截止區）', 14)
R_ROW = 4                     # 取樣電阻值（可改，所有換算欄跟著動）
HDR = 6                       # 欄組標題列
SUB = 7                       # 欄名列
R0 = 8                        # 第一列資料
RN = R0 + len(VDS_ROWS) - 1
RCELL1 = '$B$4'

ws1.cell(row=R_ROW, column=1, value='源極取樣電阻 R1 (Ω)').font = BOLD
put(ws1, R_ROW, 2, R1_EXP1, font=BLUE, fmt='0')
ws1.cell(row=R_ROW, column=3,
         value='← 講義指定 1 kΩ；改這一格，所有 I_D 與元件電壓欄會跟著變').font = NOTE

col_head(ws1, SUB, 1, 'V_DS 設定 (V)', width=14)
ws1.cell(row=HDR, column=1).border = BOX
for g, vgs in enumerate(VGS_GROUPS):
    c0 = 2 + g * 5
    group_head(ws1, HDR, c0, 5, f'V_GS 設定 = {vgs:g} V')
    col_head(ws1, SUB, c0, '實測 V_R1 (V)', width=13)
    col_head(ws1, SUB, c0 + 1, 'I_D = V_R1/R1 (mA)', width=15)
    col_head(ws1, SUB, c0 + 2, 'V_GS 元件 (V)', width=13)
    col_head(ws1, SUB, c0 + 3, 'V_DS 元件 (V)', width=13)
    col_head(ws1, SUB, c0 + 4, 'LTspice I_D (mA)', width=15)

for i, vds in enumerate(VDS_ROWS):
    r = R0 + i
    put(ws1, r, 1, vds, fmt='0.0', align='center')
    for g, vgs in enumerate(VGS_GROUPS):
        c0 = 2 + g * 5
        put(ws1, r, c0, None, fill=FILL_IN, fmt='0.000')
        vr1 = f'{get_column_letter(c0)}{r}'
        put(ws1, r, c0 + 1, f'=IF({vr1}="","",{vr1}/{RCELL1}*1000)', fmt='0.000')
        put(ws1, r, c0 + 2, f'=IF({vr1}="","",{vgs}-{vr1})', fmt='0.000')
        put(ws1, r, c0 + 3, f'=IF({vr1}="","",$A{r}-{vr1})', fmt='0.000')
        put(ws1, r, c0 + 4, round(PRED1[g][i], 4), font=BLUE, fmt='0.000')

rn = RN + 2
ws1.cell(row=rn, column=1,
         value='助教只要求至少一條導通的（V_GS = 2／4／6 V 擇一以上）；'
               'V_GS = 0 V 那組建議還是量一次，當作「截止區電流為零」的證據。').font = NOTE
ws1.cell(row=rn + 1, column=1,
         value='V_GS = 6 V 在低 V_DS 時每加 1 V 電流約多 1 mA，那是 1 kΩ 的斜率、不是元件的三極管電阻。').font = NOTE
ws1.cell(row=rn + 2, column=1,
         value='只要量 R1 的壓降這一個電壓就夠：汲極直接接在 CH2 上，V_D 必定等於電源設定值，'
               '所以元件真正的 V_GS 與 V_DS 都可以由 V_R1 減出來（右邊兩欄已自動算好）。').font = NOTE
ws1.cell(row=rn + 3, column=1,
         value='報告要畫「元件端」的特性曲線時，橫軸用 V_DS 元件 那一欄，不是最左邊的電源設定值。').font = NOTE

ch1 = ScatterChart()
ch1.title = 'I_D - V_DS 輸出特性曲線'
ch1.x_axis.title = 'V_DS 電源設定 (V)'
ch1.y_axis.title = 'I_D (mA)'
ch1.height, ch1.width = 10, 18
ch1.style = 2
xs1 = Reference(ws1, min_col=1, min_row=R0, max_row=RN)
for g, vgs in enumerate(VGS_GROUPS):
    c0 = 2 + g * 5
    ch1.series.append(dot_series(Reference(ws1, min_col=c0 + 1, min_row=R0, max_row=RN),
                                 xs1, f'實測 V_GS={vgs:g}V'))
    ch1.series.append(line_series(Reference(ws1, min_col=c0 + 4, min_row=R0, max_row=RN),
                                  xs1, f'LTspice V_GS={vgs:g}V'))
ws1.add_chart(ch1, f'A{rn + 5}')
ws1.sheet_view.showGridLines = False


# ── 實驗二：轉移特性曲線 ─────────────────────────────────────────────────────
ws2 = wb.create_sheet('實驗二 轉移曲線')
title_block(ws2, '實驗二　轉移特性曲線 I_D-V_GS 與臨界電壓（講義第 8–9 頁）',
            'R1 = 100 Ω 串在汲極、源極接地．填汲極對地電壓 V_D，I_D 自動換算成 (V_DS − V_D)/100 Ω', 14)
RCELL2 = '$B$4'
ws2.cell(row=R_ROW, column=1, value='汲極取樣電阻 R1 (Ω)').font = BOLD
put(ws2, R_ROW, 2, R1_EXP2, font=BLUE, fmt='0')
ws2.cell(row=R_ROW, column=3,
         value='← 講義指定 100 Ω；改這一格，所有 I_D 換算欄會跟著變').font = NOTE

col_head(ws2, SUB, 1, 'V_GS 設定 (V)', width=14)
ws2.cell(row=HDR, column=1).border = BOX
RN2 = R0 + len(VGS_ROWS) - 1
for g, vds in enumerate(VDS_GROUPS):
    c0 = 2 + g * 4
    group_head(ws2, HDR, c0, 4, f'V_DS 設定 = {vds:g} V')
    col_head(ws2, SUB, c0, '實測 V_D (V)', width=13)
    col_head(ws2, SUB, c0 + 1, '實測 V_R1 (V)', width=13)
    col_head(ws2, SUB, c0 + 2, 'I_D (mA)', width=12)
    col_head(ws2, SUB, c0 + 3, 'LTspice (mA)', width=13)

for i, vgs in enumerate(VGS_ROWS):
    r = R0 + i
    put(ws2, r, 1, vgs, fmt='0.0', align='center')
    for g, vds in enumerate(VDS_GROUPS):
        c0 = 2 + g * 4
        put(ws2, r, c0, None, fill=FILL_IN, fmt='0.000')
        put(ws2, r, c0 + 1, None, fill=FILL_IN, fmt='0.000')
        vd = f'{get_column_letter(c0)}{r}'
        vr1 = f'{get_column_letter(c0 + 1)}{r}'
        # lab2 的做法：優先用直接量到的 R1 壓降，沒填時才用 (V_DS − V_D) 減出來
        put(ws2, r, c0 + 2,
            f'=IF({vr1}<>"",{vr1}/{RCELL2}*1000,IF({vd}<>"",({vds}-{vd})/{RCELL2}*1000,""))',
            fmt='0.000')
        put(ws2, r, c0 + 3, round(PRED2[g][i], 4), font=BLUE, fmt='0.000')

rc = RN2 + 2
ws2.cell(row=rc, column=1,
         value='兩欄電壓量其中一個就夠，兩個都量是為了當場對帳：V_D + V_R1 應該等於該組的 V_DS 設定值，'
               '對不起來就是有一次量錯。I_D 欄優先用 V_R1，沒填時自動改用 (V_DS − V_D)/R1。').font = NOTE
ws2.cell(row=rc + 1, column=1,
         value='源極已接地，所以 V_D 就是元件真正的 V_DS、電源設定的 V_GS 就是元件真正的 V_GS，不用再減。').font = NOTE

# 加密掃描：粗掃看出在哪裡抬頭之後，回到那一段用 0.1～0.2 V 一階重掃，
# 才有兩個落在平方律直線上的點可以外插 V_th。級距當天自己填。
rf = RN2 + 5
ws2.cell(row=rf, column=1, value='加密掃描（找 V_th 用）').font = BOLD
ws2.cell(row=rf, column=3, value='選用的 V_DS (V) →').font = NOTE
put(ws2, rf, 4, None, fill=FILL_IN, fmt='0.0')
rf += 1
for j, t in enumerate(['V_GS 設定 (V)', 'V_D 量測 (V)', 'I_D (mA)', '√I_D'], start=1):
    col_head(ws2, rf, j, t, width=14 if j == 1 else 13)
F0 = rf + 1
FN = F0 + FINE_N - 1
vds_cell = f'$D${rf - 1}'
for i in range(FINE_N):
    r = F0 + i
    put(ws2, r, 1, None, fill=FILL_IN, fmt='0.00')
    put(ws2, r, 2, None, fill=FILL_IN, fmt='0.000')
    put(ws2, r, 3,
        f'=IF(OR(A{r}="",B{r}="",{vds_cell}=""),"",({vds_cell}-B{r})/{RCELL2}*1000)', fmt='0.000')
    put(ws2, r, 4, f'=IF(C{r}="","",SQRT(C{r}))', fmt='0.000')
ws2.cell(row=FN + 1, column=1,
         value='做法：先看粗掃表在哪一格之間跳起來，回到那一段用 0.1～0.2 V 一階重掃十點；'
               '√I_D 欄若呈一條直線就代表那幾點在平方律上，可以拿來外插。').font = NOTE

# 臨界電壓：平方根外插（挑兩個在抬頭段、電流 1～5 mA 之間的點）
rv = FN + 3
ws2.cell(row=rv, column=1, value='臨界電壓 V_th（平方根外插）').font = BOLD
rv += 1
for j, t in enumerate(['', '量測點 1', '量測點 2', '外插結果'], start=1):
    col_head(ws2, rv, j, t, width=16 if j == 1 else 13)
labels = [('V_GS (V)', FILL_IN, '0.00'), ('I_D (mA)', FILL_IN, '0.000')]
for k, (lab, fill, fmt) in enumerate(labels):
    r = rv + 1 + k
    put(ws2, r, 1, lab)
    put(ws2, r, 2, None, fill=fill, fmt=fmt)
    put(ws2, r, 3, None, fill=fill, fmt=fmt)
g1, i1, g2, i2 = f'B{rv + 1}', f'B{rv + 2}', f'C{rv + 1}', f'C{rv + 2}'
put(ws2, rv + 1, 4,
    f'=IF(OR({g1}="",{i1}="",{g2}="",{i2}=""),"",'
    f'{g1}-({g2}-{g1})/(SQRT({i2})-SQRT({i1}))*SQRT({i1}))', fmt='0.000')
ws2.cell(row=rv + 1, column=5, value='← V_th (V)').font = BOLD
put(ws2, rv + 2, 4, 2.14, font=BLUE, fmt='0.000')
ws2.cell(row=rv + 2, column=5, value='← LTspice（BS170）的 V_to').font = NOTE
# 防呆：兩點若取在被 R1 限流的平台上，外插會得到負值或大於量測點的荒謬結果
put(ws2, rv + 3, 4,
    f'=IF(D{rv + 1}="","",IF(OR(D{rv + 1}<0,D{rv + 1}>MIN({g1},{g2})),'
    f'"兩點取在平台上了，重挑","合理"))')
ws2.cell(row=rv + 3, column=5, value='← 合理性檢查').font = NOTE
ws2.cell(row=rv + 4, column=1,
         value='兩個點都要取自上面的加密掃描、且落在抬頭段（建議 I_D 在 1～5 mA 之間）。'
               '取到被 R1 限流的平台（I_D 接近 V_DS/100 Ω）會算出負的 V_th。').font = NOTE
ws2.cell(row=rv + 5, column=1,
         value='三組的平台高度約等於 V_DS / 100 Ω，那是 R1 在限流、不是元件飽和；要讀的資訊全在抬頭那一段。').font = NOTE

ch2 = ScatterChart()
ch2.title = 'I_D - V_GS 轉移特性曲線'
ch2.x_axis.title = 'V_GS (V)'
ch2.y_axis.title = 'I_D (mA)'
ch2.height, ch2.width = 10, 18
ch2.style = 2
xs2 = Reference(ws2, min_col=1, min_row=R0, max_row=RN2)
for g, vds in enumerate(VDS_GROUPS):
    c0 = 2 + g * 4
    ch2.series.append(dot_series(Reference(ws2, min_col=c0 + 2, min_row=R0, max_row=RN2),
                                 xs2, f'實測 V_DS={vds:g}V'))
    ch2.series.append(line_series(Reference(ws2, min_col=c0 + 3, min_row=R0, max_row=RN2),
                                  xs2, f'LTspice V_DS={vds:g}V'))
ws2.add_chart(ch2, f'A{rv + 7}')

ch3 = ScatterChart()
ch3.title = '√I_D - V_GS（加密掃描；直線段外插到 x 軸就是 V_th）'
ch3.x_axis.title = 'V_GS (V)'
ch3.y_axis.title = '√I_D (√mA)'
ch3.height, ch3.width = 9, 18
ch3.style = 2
ch3.series.append(dot_series(Reference(ws2, min_col=4, min_row=F0, max_row=FN),
                             Reference(ws2, min_col=1, min_row=F0, max_row=FN), '加密掃描'))
ws2.add_chart(ch3, f'L{rv + 7}')
ws2.sheet_view.showGridLines = False


# ── 交報告用：照講義第 7／9 頁的版面 ─────────────────────────────────────────
ws3 = wb.create_sheet('交報告用')
title_block(ws3, '交報告用表格（版面照 App_Lab3.pdf 第 7 頁與第 9 頁）',
            '數字由前兩張表帶過來；空白代表當天沒量。直接整塊複製貼進報告', 8)

r = 4
ws3.cell(row=r, column=1, value='1. Circuit schematic (characteristic curve, I_D-V_DS)').font = BOLD
r += 1
# 2×2 四個區塊，和講義排法一樣
for gi, vgs in enumerate(VGS_GROUPS):
    base_r = r + (gi // 2) * (len(VDS_ROWS) + 3)
    base_c = 1 + (gi % 2) * 3
    group_head(ws3, base_r, base_c, 2, f'V_GS = {vgs:g} V')
    col_head(ws3, base_r + 1, base_c, 'V_DS (experimental setting)', width=26)
    col_head(ws3, base_r + 1, base_c + 1, 'I_D (measure)', width=16)
    src_c = get_column_letter(2 + gi * 5 + 1)
    for i, vds in enumerate(VDS_ROWS):
        rr = base_r + 2 + i
        put(ws3, rr, base_c, f'{vds:g} V', align='center')
        put(ws3, rr, base_c + 1,
            f"=IF('實驗一 特性曲線'!{src_c}{R0 + i}=\"\",\"\","
            f"TEXT('實驗一 特性曲線'!{src_c}{R0 + i},\"0.000\")&\" mA\")", align='center')

r = r + 2 * (len(VDS_ROWS) + 3) + 1
ws3.cell(row=r, column=1, value='2. Circuit schematic (transfer curve, I_D-V_GS)').font = BOLD
r += 1
for gi, vds in enumerate(VDS_GROUPS):
    base_r = r + gi * (len(VGS_ROWS) + 3)
    group_head(ws3, base_r, 1, 2, f'V_DS = {vds:g} V')
    col_head(ws3, base_r + 1, 1, 'V_GS (experimental setting)', width=26)
    col_head(ws3, base_r + 1, 2, 'I_D (measure)', width=16)
    src_c = get_column_letter(2 + gi * 4 + 2)
    for i, vgs in enumerate(VGS_ROWS):
        rr = base_r + 2 + i
        put(ws3, rr, 1, f'{vgs:g} V', align='center')
        put(ws3, rr, 2,
            f"=IF('實驗二 轉移曲線'!{src_c}{R0 + i}=\"\",\"\","
            f"TEXT('實驗二 轉移曲線'!{src_c}{R0 + i},\"0.000\")&\" mA\")", align='center')
ws3.sheet_view.showGridLines = False

# 公式格由 Excel / LibreOffice 開啟時重新計算，避免預覽器沿用舊快取值。
wb.calculation.calcMode = 'auto'
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True

os.makedirs(os.path.dirname(OUT), exist_ok=True)
if os.path.exists(OUT) and '--force' not in sys.argv:
    # 作者會直接在這個檔裡填數據、調格式，所以預設不覆蓋；要重產範本加 --force。
    sys.exit(f'{OUT} 已存在，未覆蓋（裡面可能有作者填的數據與格式）。確定要重產請加 --force。')
wb.save(OUT)
print('wrote', OUT)
